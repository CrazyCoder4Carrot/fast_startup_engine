"""Build docs/slides.pptx from docs/slides.html as editable slides, for import into Google Slides
(File -> Import slides). Text, lists and tables become native text boxes and tables; cards become
shapes; only pure graphics (charts, timelines, diagrams, legends) are embedded as images. Every
element is placed from its position in the rendered HTML, so the layout matches.

    uv run python docs/make_slides.py
    uv run --with playwright --with python-pptx python docs/make_pptx.py   # needs: playwright install chromium
"""

import os
import re

from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
SRC, OUT = os.path.join(HERE, "slides.html"), os.path.join(HERE, "slides.pptx")
VW, VH = 1600, 900                      # render size (16:9); positions are fractions of the slide
SW, SH = Emu(12192000), Emu(6858000)    # 13.333 in x 7.5 in
PX_TO_PT = 13.333 * 72 / VW             # CSS px at 1600 wide -> points
FONT, MONO = "Inter", "Roboto Mono"
SHOTS = "/tmp/pptx_parts"

# Walk each slide's DOM and describe it as positioned blocks. Graphics become "image" blocks
# (screenshotted later by data-exp id); boxes keep their own text; text leaves carry styled runs.
EXTRACT = r"""
(ix) => {
  const sec = document.querySelector('section.active'), R = sec.getBoundingClientRect();
  const box = (e) => { const r = e.getBoundingClientRect();
    return {x: (r.left - R.left) / R.width, y: (r.top - R.top) / R.height, w: r.width / R.width, h: r.height / R.height}; };
  const cs = (e) => getComputedStyle(e);
  const IMG = 'svg, .tl, .bs, .mem, .chart, .legend, .frames';
  const BOX = '.panel, .step, .fbox, .box, .ly, .vseg, .guard, .pc-box, .hero > div, pre.code';
  const BG = '.lane';
  const TEXT = 'h1, h2, h4, h5, p, ul, ol, .kicker, .num, .flow-tag, .farr, .arr, .down, .vas-title, .bs-axis, .footer span, .badge';
  let n = 0; const blocks = [];
  // Styled runs of an element's text: bold / mono / colour per text node.
  const runs = (el) => { const out = [];
    const walk = (node, inherited) => {
      for (const c of node.childNodes) {
        if (c.nodeType === 3) { const t = c.textContent.replace(/\s+/g, ' ');
          if (t.trim() || (t === ' ' && out.length)) out.push({t, ...inherited}); }
        else if (c.nodeType === 1) {
          if (c.matches('br')) { out.push({t: '\n', ...inherited}); continue; }
          const s = cs(c); if (s.display === 'none') continue;
          if (c.matches(IMG) || s.position === 'absolute') continue;  // emitted separately
          const st = {b: parseInt(s.fontWeight) >= 600, mono: /mono|menlo/i.test(s.fontFamily), color: s.color, size: parseFloat(s.fontSize),
                      upper: s.textTransform === 'uppercase'};
          const block = ['block', 'list-item', 'flex', 'grid'].includes(s.display) && out.length;
          if (block) out.push({t: '\n', ...inherited});
          walk(c, st);
        }
      }
    };
    const s = cs(el);
    walk(el, {b: parseInt(s.fontWeight) >= 600, mono: /mono|menlo/i.test(s.fontFamily), color: s.color, size: parseFloat(s.fontSize),
              upper: s.textTransform === 'uppercase'});
    return out.map(r => ({...r, t: r.upper ? r.t.toUpperCase() : r.t}));
  };
  const listItems = (el) => [...el.children].filter(li => li.tagName === 'LI').map((li, k) =>
    ({runs: runs(li), num: el.tagName === 'OL' ? k + 1 : null}));
  const visit = (el) => {
    const s = cs(el); if (s.display === 'none' || s.visibility === 'hidden') return;
    const b = box(el); if (b.w <= 0 || b.h <= 0) return;
    if (el.matches(IMG)) { el.dataset.exp = `${ix}-${n}`; blocks.push({type: 'image', id: `${ix}-${n++}`, ...b}); return; }
    if (el.matches('table')) {
      const rows = [...el.rows].map(tr => [...tr.cells].map(td => ({runs: runs(td), span: td.colSpan, head: td.tagName === 'TH',
        bg: cs(td).backgroundColor, color: cs(td).color, bold: parseInt(cs(td).fontWeight) >= 600, size: parseFloat(cs(td).fontSize)})));
      blocks.push({type: 'table', rows, ...b}); return; }
    if (el.matches(BOX)) {
      blocks.push({type: 'box', bg: s.backgroundColor, border: s.borderLeftColor, borderW: parseFloat(s.borderLeftWidth),
        edge: s.borderTopColor, radius: parseFloat(s.borderTopLeftRadius), pad: [parseFloat(s.paddingTop), parseFloat(s.paddingLeft)],
        runs: runs(el), lists: [...el.querySelectorAll('ol, ul')].map(listItems), ...b});
      for (const g of el.querySelectorAll(IMG)) {
        g.dataset.exp = `${ix}-${n}`; blocks.push({type: 'image', id: `${ix}-${n++}`, ...box(g)}); }
      for (const a of el.querySelectorAll('*')) if (cs(a).position === 'absolute' && a.textContent.trim())
        blocks.push({type: 'text', runs: runs(a), align: 'left', lh: 1.2, ...box(a)});
      return; }
    if (el.matches(BG)) blocks.push({type: 'box', bg: s.backgroundColor, border: 'rgba(0,0,0,0)', borderW: 0, edge: 'rgba(0,0,0,0)',
      radius: parseFloat(s.borderTopLeftRadius), pad: [0, 0], runs: [], lists: [], ...b});
    if (el.matches(TEXT)) {
      if (el.matches('ul, ol')) blocks.push({type: 'list', items: listItems(el), align: s.textAlign, ...b});
      else blocks.push({type: 'text', runs: runs(el), align: el.matches('.arr, .farr, .down') ? 'center' : s.textAlign,
        middle: el.matches('.arr, .farr, .down'), lh: parseFloat(s.lineHeight) / parseFloat(s.fontSize), ...b});
      return; }
    for (const c of el.children) visit(c);
  };
  for (const c of sec.children) visit(c);
  return {bg: cs(sec).backgroundColor, bgImage: cs(sec).backgroundImage, cls: sec.className, blocks};
}
"""


def rgb(css: str):
    """CSS rgb()/rgba() -> (RGBColor, alpha) or (None, 0)."""
    m = re.match(r"rgba?\(([\d.]+),\s*([\d.]+),\s*([\d.]+)(?:,\s*([\d.]+))?\)", css or "")
    if not m:
        return None, 0.0
    a = float(m.group(4)) if m.group(4) is not None else 1.0
    return RGBColor(*(int(float(m.group(i))) for i in (1, 2, 3))), a


def emu(b: dict):
    return Emu(int(b["x"] * SW)), Emu(int(b["y"] * SH)), Emu(max(int(b["w"] * SW), 1)), Emu(max(int(b["h"] * SH), 1))


def fill_runs(p, runs, default_size=None):
    """Add styled runs to a paragraph; '\\n' runs start new paragraphs (returns the last one)."""
    tf = p._parent
    for r in runs:
        parts = r["t"].split("\n")
        for k, part in enumerate(parts):
            if k:
                p = tf.add_paragraph()
                p.alignment = tf.paragraphs[0].alignment
            if not part:
                continue
            run = p.add_run()
            run.text = part
            f = run.font
            f.name = MONO if r.get("mono") else FONT
            f.bold = bool(r.get("b"))
            f.size = Pt(round((r.get("size") or default_size or 16) * PX_TO_PT, 1))
            c, a = rgb(r.get("color"))
            if c is not None and a > 0:
                f.color.rgb = c
    return p


def textbox(slide, b, runs, align="left", pad=(0, 0), anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(*emu(b))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_top = tf.margin_bottom = Emu(int(pad[0] / VH * SH))
    tf.margin_left = tf.margin_right = Emu(int(pad[1] / VW * SW))
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = {"right": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER, "end": PP_ALIGN.RIGHT}.get(align, PP_ALIGN.LEFT)
    fill_runs(p, runs)
    return tb


def add_box(slide, b):
    x, y, w, h = emu(b)
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if b["radius"] > 0 else MSO_SHAPE.RECTANGLE, x, y, w, h)
    if b["radius"] > 0:
        shape.adjustments[0] = min(0.5, b["radius"] / max(b["h"] * VH, 1))
    c, a = rgb(b["bg"])
    if c is not None and a > 0.05:
        shape.fill.solid(); shape.fill.fore_color.rgb = c
    else:
        shape.fill.background()
    ec, ea = rgb(b["edge"])
    if ec is not None and ea > 0.05:
        shape.line.color.rgb = ec; shape.line.width = Pt(0.75)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    if b["borderW"] >= 2:  # coloured left edge of a card
        bc, ba = rgb(b["border"])
        if bc is not None and ba > 0.05:
            bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Emu(int(b["borderW"] / VW * SW)), h)
            bar.fill.solid(); bar.fill.fore_color.rgb = bc; bar.line.fill.background(); bar.shadow.inherit = False
    if b["runs"]:
        tf = shape.text_frame
        tf.word_wrap = True
        tf.auto_size = None
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_top = tf.margin_bottom = Emu(int(b["pad"][0] / VH * SH))
        tf.margin_left = tf.margin_right = Emu(int(max(b["pad"][1], 4) / VW * SW))
        fill_runs(tf.paragraphs[0], b["runs"])


def add_list(slide, b):
    tb = slide.shapes.add_textbox(*emu(b))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for it in b["items"]:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(6)
        runs = it["runs"]
        if it["num"] is not None:
            runs = [{**runs[0], "t": f"{it['num']}. ", "b": False}] + runs if runs else runs
        elif runs and runs[0]["t"].strip() not in ("+", "−"):
            runs = [{**runs[0], "t": "•  ", "b": False}] + runs
        fill_runs(p, runs)


def add_table(slide, b):
    rows = b["rows"]
    ncols = max(sum(c["span"] for c in r) for r in rows)
    x, y, w, h = emu(b)
    gt = slide.shapes.add_table(len(rows), ncols, x, y, w, h).table
    gt.first_row = False
    gt.horz_banding = False
    for i, r in enumerate(rows):
        j = 0
        for c in r:
            cell = gt.cell(i, j)
            if c["span"] > 1:
                cell.merge(gt.cell(i, j + c["span"] - 1))
            cell.fill.background()
            bc, ba = rgb(c["bg"])
            if bc is not None and ba > 0.05:
                cell.fill.solid(); cell.fill.fore_color.rgb = bc
            tf = cell.text_frame
            tf.word_wrap = True
            cell.margin_left = cell.margin_right = Emu(int(8 / VW * SW))
            cell.margin_top = cell.margin_bottom = Emu(int(4 / VH * SH))
            runs = [{**rr, "b": rr.get("b") or c["bold"]} for rr in c["runs"]] or [{"t": "", "size": c["size"]}]
            fill_runs(tf.paragraphs[0], runs, default_size=c["size"])
            j += c["span"]


def main():
    os.makedirs(SHOTS, exist_ok=True)
    prs = Presentation()
    prs.slide_width, prs.slide_height = SW, SH
    with sync_playwright() as pw:
        br = pw.chromium.launch()
        pg = br.new_page(viewport={"width": VW, "height": VH}, device_scale_factor=2)
        pg.goto("file://" + SRC)
        pg.wait_for_timeout(800)
        n = pg.evaluate("document.querySelectorAll('section').length")
        for i in range(n):
            pg.evaluate(f"show({i})")
            pg.wait_for_timeout(120)
            d = pg.evaluate(EXTRACT, i)
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            c, a = rgb(d["bg"])
            if "cover" in d["cls"] or "divider" in d["cls"]:
                c, a = RGBColor(0x0B, 0x12, 0x20), 1.0   # the dark title background (a CSS gradient in HTML)
            if c is not None and a > 0:
                slide.background.fill.solid(); slide.background.fill.fore_color.rgb = c
            notes = []
            for blk in d["blocks"]:
                t = blk["type"]
                if t == "image":
                    path = f"{SHOTS}/{blk['id']}.png"
                    pg.locator(f'[data-exp="{blk["id"]}"]').screenshot(path=path)
                    slide.shapes.add_picture(path, *emu(blk))
                elif t == "table":
                    add_table(slide, blk)
                elif t == "box":
                    add_box(slide, blk)
                elif t == "list":
                    add_list(slide, blk)
                    notes += ["".join(r["t"] for r in it["runs"]) for it in blk["items"]]
                else:
                    textbox(slide, blk, blk["runs"], blk.get("align", "left"),
                            anchor=MSO_ANCHOR.MIDDLE if blk.get("middle") else MSO_ANCHOR.TOP)
                    notes.append("".join(r["t"] for r in blk["runs"]))
            slide.notes_slide.notes_text_frame.text = "\n".join(x for x in notes if x.strip())
        br.close()
    prs.save(OUT)
    print(f"wrote {OUT} ({n} slides, {os.path.getsize(OUT) / 1e6:.1f} MB)")


if __name__ == "__main__":
    main()
