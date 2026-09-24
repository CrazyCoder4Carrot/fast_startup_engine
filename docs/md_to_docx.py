"""Minimal Markdown -> .docx for importing the report into Google Docs."""
import re, sys, os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

src, dst = sys.argv[1], sys.argv[2]
base = os.path.dirname(os.path.abspath(src))
doc = Document()
st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10.5)

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)")
def add_runs(p, text):
    for part in INLINE.split(text):
        if not part: continue
        if part.startswith("**"):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`"):
            r = p.add_run(part[1:-1]); r.font.name = "Courier New"; r.font.size = Pt(9.5)
        elif part.startswith("*") and len(part) > 2:
            r = p.add_run(part[1:-1]); r.italic = True
        else:
            p.add_run(part)

lines = open(src).read().splitlines()
i = 0
while i < len(lines):
    ln = lines[i]
    if ln.startswith("```"):
        i += 1; code = []
        while not lines[i].startswith("```"):
            code.append(lines[i]); i += 1
        p = doc.add_paragraph(); r = p.add_run("\n".join(code))
        r.font.name = "Courier New"; r.font.size = Pt(9)
    elif m := re.match(r"^(#{1,3}) (.*)", ln):
        doc.add_heading(m.group(2), level=len(m.group(1)) - 1)
    elif m := re.match(r"^!\[(.*?)\]\((.*?)\)", ln):
        doc.add_picture(os.path.join(base, m.group(2)), width=Inches(6.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif ln.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            cells = [c.strip() for c in lines[i].strip("|").split("|")]
            if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                rows.append(cells)
            i += 1
        i -= 1
        t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"
        for r_i, row in enumerate(rows):
            for c_i, cell in enumerate(row):
                p = t.cell(r_i, c_i).paragraphs[0]
                add_runs(p, cell if r_i else f"**{cell}**" if cell else "")
        doc.add_paragraph()
    elif ln.startswith("> "):
        p = doc.add_paragraph(); add_runs(p, ln[2:])
        for r in p.runs: r.italic = True; r.font.color.rgb = RGBColor(0x52, 0x51, 0x4e)
    elif m := re.match(r"^(\s*)- (.*)", ln):
        style = "List Bullet 2" if len(m.group(1)) >= 2 else "List Bullet"
        add_runs(doc.add_paragraph(style=style), m.group(2))
    elif m := re.match(r"^(\d+)\. (.*)", ln):
        add_runs(doc.add_paragraph(style="List Number"), m.group(2))
    elif ln.strip():
        add_runs(doc.add_paragraph(), ln.strip())
    i += 1
doc.save(dst)
print("wrote", dst)
